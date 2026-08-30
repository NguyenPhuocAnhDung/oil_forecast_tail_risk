import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

src_doc = docx.Document('GUMNETHet_FAIRv3.docx')

print("=== CHECKING ALL 13 TABLES IN FAIRv3 ===")
for i, t in enumerate(src_doc.tables):
    rows = len(t.rows)
    cols = len(t.columns)
    first_cell = t.cell(0,0).text.strip()[:30]
    widths = [c.width for c in t.rows[0].cells]
    print(f"Table {i:02d} ({rows}x{cols}): widths={widths} | text={repr(first_cell)}")
